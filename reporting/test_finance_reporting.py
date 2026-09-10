from __future__ import annotations

import json
import shutil
import tempfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounting.models import (
    AccountingPeriod, FiscalYear, Fund, JournalEntry, JournalLine, JournalSubsidiaryLine,
    LedgerAccount, PostingMapping, ResponsibilityCenter,
)
from budget.models import (
    AllotmentMovement, AllotmentReleaseOrder, AppropriationAuthorization,
    AuthorizedAppropriationLine, BudgetCall, BudgetVersion, ObligationMovement,
    ObligationRequest,
)
from departments.models import Department
from finance.models import (
    FinanceConfigurationItem, FinanceConfigurationRelease, finance_tax_rule_snapshot,
)
from vouchers.models import (
    BankAdviceBatch, BankAdviceItem, DisbursementVoucher, PaymentInstrument, VoucherCase,
    voucher_tax_evidence_checksum,
)

from .datasets import build_dataset_with_evidence
from .models import (
    FinanceStatementLine, FinanceStatementMapping, ReportDefinition, ReportRun,
    ReportTemplateVersion,
)
from .presets import seed_finance_presets
from .services import create_manual_run, transition_run
from .statement_services import review_statement_mapping, submit_statement_mapping


FINANCE_REPORT_MEDIA_ROOT = tempfile.mkdtemp(prefix="grand-finance-report-tests-")


@override_settings(MEDIA_ROOT=FINANCE_REPORT_MEDIA_ROOT)
class FinanceAccountabilityReportingTests(TestCase):
    databases = {"default", "finance"}

    @classmethod
    def setUpTestData(cls):
        cls.accounting = Department.objects.create(name="Municipal Accounting Office", slug="f9-accounting")
        cls.budget = Department.objects.create(name="Municipal Budget Office", slug="f9-budget")
        cls.requesting = Department.objects.create(name="General Services Office", slug="f9-gso")
        cls.accounting_preparer = cls.employee(
            cls.accounting, "f9.accounting.preparer",
            "view_reporting_workspace", "generate_reports", "download_reports",
            "manage_report_definitions",
        )
        cls.accounting_reviewer = cls.employee(
            cls.accounting, "f9.accounting.reviewer",
            "view_reporting_workspace", "review_reports", "approve_reports",
            "download_reports", "view_department_reports",
        )
        cls.budget_preparer = cls.employee(
            cls.budget, "f9.budget.preparer",
            "view_reporting_workspace", "generate_reports", "download_reports",
        )
        cls.accounting.deptHead_or_oic = cls.accounting_reviewer
        cls.accounting.save(update_fields=("deptHead_or_oic",))
        cls.budget.deptHead_or_oic = cls.budget_preparer
        cls.budget.save(update_fields=("deptHead_or_oic",))

        owner = {"department_id": cls.accounting.pk, "department_label": cls.accounting.name}
        cls.fiscal_year = FiscalYear.objects.create(
            **owner, year=2027, label="FY 2027", starts_on=date(2027, 1, 1),
            ends_on=date(2027, 12, 31), business_date=date(2027, 3, 31),
            status=FiscalYear.ACTIVE,
        )
        cls.period = AccountingPeriod.objects.create(
            **owner, fiscal_year=2027, fiscal_year_record=cls.fiscal_year, period_number=1,
            label="First quarter", starts_on=date(2027, 1, 1), ends_on=date(2027, 3, 31),
        )
        cls.fund = Fund.objects.create(**owner, code="GF", name="General Fund")
        cls.cash = LedgerAccount.objects.create(
            **owner, code="10101010", title="Cash in bank", account_type="asset", normal_balance="debit",
        )
        cls.revenue = LedgerAccount.objects.create(
            **owner, code="40101010", title="Local revenue", account_type="revenue", normal_balance="credit",
        )
        cls.entry = JournalEntry.objects.create(
            **owner, reference="JEV-F9-0001", entry_date=date(2027, 1, 15), period=cls.period,
            fund=cls.fund, source_type="manual", description="Synthetic balanced accountability entry",
            status=JournalEntry.DRAFT, created_by_id=cls.accounting_preparer.pk,
            created_by_label=cls.accounting_preparer.username,
            posted_by_id=cls.accounting_reviewer.pk, posted_by_label=cls.accounting_reviewer.username,
            posted_at=timezone.now(),
        )
        JournalLine.objects.create(entry=cls.entry, sequence=1, account=cls.cash, debit=Decimal("1250.00"))
        JournalLine.objects.create(entry=cls.entry, sequence=2, account=cls.revenue, credit=Decimal("1250.00"))
        JournalEntry.objects.filter(pk=cls.entry.pk).update(status=JournalEntry.POSTED)
        cls.entry.refresh_from_db()

        call = BudgetCall.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name,
            fiscal_year=cls.fiscal_year, title="FY 2027 Budget Call",
            authority_reference="Synthetic reviewed budget call", instructions="Synthetic instructions",
            proposal_opens_on=date(2026, 8, 1), proposal_due_on=date(2026, 9, 30),
            status=BudgetCall.PUBLISHED, created_by_id=cls.budget_preparer.pk,
            created_by_label=cls.budget_preparer.username,
        )
        version = BudgetVersion.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name,
            budget_call=call, fiscal_year=cls.fiscal_year, kind=BudgetVersion.FINAL, version=1,
            title="FY 2027 Authorized Budget", change_explanation="Synthetic final authority",
            status=BudgetVersion.AUTHORIZED, created_by_id=cls.budget_preparer.pk,
            created_by_label=cls.budget_preparer.username,
        )
        authorization = AppropriationAuthorization.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name, version=version,
            authority_type=AppropriationAuthorization.ORDINANCE,
            ordinance_number="SYN-ORD-2027-001", ordinance_date=date(2026, 12, 15),
            effectivity_date=date(2027, 1, 1), review_status=AppropriationAuthorization.FAVORABLE,
            review_reference="Synthetic favorable review", review_date=date(2026, 12, 28),
            evidence_reference="Synthetic retained ordinance and schedule",
            signed_control_total=Decimal("100000.00"), status=AppropriationAuthorization.DRAFT,
            snapshot_checksum="a" * 64, created_by_id=cls.budget_preparer.pk,
            created_by_label=cls.budget_preparer.username,
        )
        cls.appropriation_line = AuthorizedAppropriationLine.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name,
            authorization=authorization, source_line_id=9001, fund_code="GF",
            responsibility_center_code="GSO", program_code="GSO-OPS",
            funding_source_code="LOCAL", account_code="5-02-03",
            expense_class="MOOE", appropriation_type="new",
            particulars="Synthetic office operations", amount=Decimal("100000.00"),
        )
        AppropriationAuthorization.objects.filter(pk=authorization.pk).update(status=AppropriationAuthorization.AUTHORIZED)
        authorization.refresh_from_db()
        order = AllotmentReleaseOrder.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name,
            authorization=authorization, fiscal_year=cls.fiscal_year,
            order_number="ARO-F9-0001", kind=AllotmentReleaseOrder.INITIAL,
            release_date=date(2027, 1, 5), effective_date=date(2027, 1, 5),
            authority_reference="Synthetic allotment authority",
            evidence_reference="Synthetic signed allotment schedule", purpose="Quarterly operations",
            signed_control_total=Decimal("80000.00"), status=AllotmentReleaseOrder.POSTED,
            snapshot_checksum="b" * 64, created_by_id=cls.budget_preparer.pk,
            created_by_label=cls.budget_preparer.username,
        )
        AllotmentMovement.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name,
            order=order, source_line_id=8001, appropriation_line=cls.appropriation_line,
            movement_type="release", amount=Decimal("80000.00"),
            release_effect=Decimal("80000.00"), hold_effect=Decimal("0.00"),
            effective_date=date(2027, 1, 5), order_number_snapshot=order.order_number,
            authority_reference_snapshot=order.authority_reference,
        )
        obligation = ObligationRequest.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name,
            authorization=authorization, fiscal_year=cls.fiscal_year,
            requesting_department_id=cls.requesting.pk, requesting_department_label=cls.requesting.name,
            kind=ObligationRequest.ORIGINAL, form_type=ObligationRequest.OBR,
            request_reference="REQ-F9-0001", obligation_number="OBR-F9-0001",
            obligation_date=date(2027, 2, 10), claimant_payee="Synthetic Supplier",
            particulars="Synthetic accountable obligation", evidence_reference="Synthetic retained support",
            signed_control_total=Decimal("50000.00"), status=ObligationRequest.CERTIFIED,
            snapshot_checksum="c" * 64, created_by_id=cls.budget_preparer.pk,
            created_by_label=cls.budget_preparer.username,
        )
        ObligationMovement.objects.create(
            department_id=cls.budget.pk, department_label=cls.budget.name,
            request=obligation, source_line_id=7001, appropriation_line=cls.appropriation_line,
            movement_type="obligate", amount=Decimal("50000.00"),
            obligation_effect=Decimal("50000.00"), effective_date=date(2027, 2, 10),
            obligation_number_snapshot=obligation.obligation_number,
            requesting_department_snapshot=cls.requesting.name,
            claimant_payee_snapshot=obligation.claimant_payee,
            particulars_snapshot=obligation.particulars,
        )
        seed_finance_presets()
        cls.accounting_definition = ReportDefinition.objects.get(
            department=cls.accounting, dataset_key="finance_posted_trial_balance",
        )
        cls.budget_definition = ReportDefinition.objects.get(
            department=cls.budget, dataset_key="finance_budget_accountability",
        )

    @classmethod
    def employee(cls, department, username, *permissions):
        user = get_user_model().objects.create_user(
            username=username, email=f"{username}@example.test", password="finance-report-test",
        )
        user.employeeprofile.assigned_department = department
        user.employeeprofile.save(update_fields=("assigned_department",))
        user.user_permissions.add(*Permission.objects.filter(
            content_type__app_label="reporting", codename__in=permissions,
        ))
        return user

    def generate_accounting(self, actor=None):
        return create_manual_run(
            self.accounting_definition, self.accounting_definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, actor or self.accounting_preparer,
        )

    def test_governed_tax_detail_and_summary_reconcile_rule_formula_ledger_and_reversal(self):
        release = FinanceConfigurationRelease.objects.create(
            department=self.accounting, code="f9-tax", version=1,
            title="Synthetic governed tax setup", fiscal_year=2027,
            status="active", effective_from=date(2027, 1, 1),
            created_by=self.accounting_preparer, activated_by=self.accounting_reviewer,
            activated_at=timezone.now(),
        )
        tax_item = FinanceConfigurationItem.objects.create(
            department=self.accounting, release=release, category="tax_rule", code="ewt-1",
            version=1, label="Synthetic expanded withholding", description="Synthetic UAT rule",
            configuration={
                "reporting_enabled": True, "tax_family": "expanded_income", "atc": "WI158",
                "rate_percent": "1", "tax_base_label": "Reviewed gross income payment",
                "return_form_code": "1601-EQ", "certificate_form_code": "2307",
                "reporting_basis": "accounting_posting", "rounding_mode": "half_up",
                "requires_tax_identifier": True,
                "authority_reference": "Synthetic locally reviewed BIR rule",
                "applicability_status": "locally_confirmed",
                "local_acceptance_note": "Synthetic Accounting acceptance evidence",
            },
            status="active", effective_from=date(2027, 1, 1), created_by=self.accounting_preparer,
        )
        rule_snapshot, rule_checksum = finance_tax_rule_snapshot(tax_item)
        tax_snapshot = {
            **rule_snapshot, "tax_base": "1000.00", "tax_withheld": "10.00",
            "tax_rule_checksum": rule_checksum, "payee_name": "Synthetic Supplier",
            "payee_tax_identifier": "000-000-001-00000", "voucher_date": "2027-02-15",
            "voucher_number": "DV-TAX-001", "case_reference": "CASE-TAX-001",
            "case_public_id": "11111111-1111-1111-1111-111111111111",
        }
        tax_snapshot["tax_evidence_checksum"] = voucher_tax_evidence_checksum(
            voucher=SimpleNamespace(dv_number="DV-TAX-001", voucher_date=date(2027, 2, 15)),
            tax_rule_checksum=rule_checksum, tax_base=Decimal("1000.00"), amount=Decimal("10.00"),
            payee_name="Synthetic Supplier", payee_tax_identifier="000-000-001-00000",
        )
        withholding = LedgerAccount.objects.create(
            department_id=self.accounting.pk, department_label=self.accounting.name,
            code="2-02-TAX-F9", title="Synthetic governed tax payable",
            account_type="liability", normal_balance="credit",
        )
        tax_entry = JournalEntry.objects.create(
            department_id=self.accounting.pk, department_label=self.accounting.name,
            reference="JEV-TAX-F9-001", entry_date=date(2027, 2, 15), period=self.period,
            fund=self.fund, source_type="voucher", source_reference="tax-source-1",
            description="Synthetic governed withholding", status=JournalEntry.DRAFT,
            created_by_id=self.accounting_preparer.pk, created_by_label=self.accounting_preparer.username,
            posted_by_id=self.accounting_reviewer.pk, posted_by_label=self.accounting_reviewer.username,
            posted_at=timezone.now(),
        )
        debit = JournalLine.objects.create(
            entry=tax_entry, sequence=1, account=self.cash, debit=Decimal("10.00"),
        )
        credit = JournalLine.objects.create(
            entry=tax_entry, sequence=2, account=withholding, credit=Decimal("10.00"),
        )
        tax_detail = JournalSubsidiaryLine.objects.create(
            entry=tax_entry, journal_line=credit, category=JournalSubsidiaryLine.WITHHOLDING,
            reference_key="ewt-1", reference_label="Synthetic expanded withholding",
            source_code="ewt-1", source_reference="tax-source-1", credit=Decimal("10.00"),
            source_snapshot={"tax_reporting": tax_snapshot},
        )
        JournalEntry.objects.filter(pk=tax_entry.pk).update(status=JournalEntry.POSTED)
        tax_entry.refresh_from_db()
        self.assertEqual(debit.debit, Decimal("10.00"))

        detail_definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_governed_tax_withholding_detail",
        )
        _adapter, rows, totals, evidence = build_dataset_with_evidence(
            detail_definition, date(2027, 1, 1), date(2027, 3, 31), {},
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["payee_tax_identifier"], "000-000-001-00000")
        self.assertEqual(rows[0]["tax_withheld"], Decimal("10.00"))
        self.assertEqual(evidence["control_status"], ReportRun.CONTROL_RECONCILED)
        self.assertEqual(evidence["control_totals"]["ledger_difference"], Decimal("0.00"))
        self.assertEqual(totals["tax_withheld"], Decimal("10.00"))
        run = create_manual_run(
            detail_definition, detail_definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(run.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(run.dataset_snapshot["rows"][0]["payee_tax_identifier"], "000-000-001-00000")
        self.assertEqual(len(run.dataset_checksum), 64)
        self.assertTrue(run.output_file.name.endswith(".xlsx"))
        # Synthetic five-digit branch-code compatibility; no official tax rule
        # or filing authority is inferred from preserving the identifier.
        from openpyxl import load_workbook
        workbook = load_workbook(run.output_file.path, read_only=True, data_only=True)
        try:
            self.assertTrue(any(
                cell == "000-000-001-00000"
                for sheet in workbook for row in sheet.iter_rows(values_only=True) for cell in row
            ))
        finally:
            workbook.close()

        reversal = JournalEntry.objects.create(
            department_id=self.accounting.pk, department_label=self.accounting.name,
            reference="JEV-TAX-F9-REV", entry_date=date(2027, 3, 1), period=self.period,
            fund=self.fund, source_type="reversal", source_reference="tax-reversal-1",
            reversal_of=tax_entry, reversal_reason="Synthetic correction",
            description="Reverse synthetic governed withholding", status=JournalEntry.DRAFT,
            created_by_id=self.accounting_preparer.pk, created_by_label=self.accounting_preparer.username,
            posted_by_id=self.accounting_reviewer.pk, posted_by_label=self.accounting_reviewer.username,
            posted_at=timezone.now(),
        )
        reversal_debit = JournalLine.objects.create(
            entry=reversal, sequence=1, account=withholding, debit=Decimal("10.00"),
        )
        JournalLine.objects.create(entry=reversal, sequence=2, account=self.cash, credit=Decimal("10.00"))
        JournalSubsidiaryLine.objects.create(
            entry=reversal, journal_line=reversal_debit, category=JournalSubsidiaryLine.WITHHOLDING,
            reference_key="ewt-1", reference_label="Synthetic expanded withholding",
            source_code="ewt-1", source_reference="tax-reversal-1", debit=Decimal("10.00"),
            source_snapshot={"reversal_of_subsidiary_line": tax_detail.pk, "tax_reporting": tax_snapshot},
        )
        JournalEntry.objects.filter(pk=reversal.pk).update(status=JournalEntry.POSTED)
        reversal.refresh_from_db()
        summary_definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_governed_tax_return_summary",
        )
        _adapter, summary_rows, summary_totals, summary_evidence = build_dataset_with_evidence(
            summary_definition, date(2027, 1, 1), date(2027, 3, 31), {},
        )
        self.assertEqual(summary_rows[0]["line_count"], 2)
        self.assertEqual(summary_rows[0]["tax_withheld"], Decimal("0.00"))
        self.assertEqual(summary_totals["tax_withheld"], Decimal("0.00"))
        self.assertEqual(summary_evidence["control_status"], ReportRun.CONTROL_RECONCILED)

        broken = dict(tax_snapshot)
        broken["rate_percent"] = "2"
        JournalSubsidiaryLine.objects.filter(pk=tax_detail.pk).update(
            source_snapshot={"tax_reporting": broken},
        )
        _adapter, _rows, _totals, broken_evidence = build_dataset_with_evidence(
            detail_definition, date(2027, 1, 1), date(2027, 3, 31), {},
        )
        self.assertEqual(broken_evidence["control_status"], ReportRun.CONTROL_EXCEPTION)
        self.assertGreater(broken_evidence["control_totals"]["invalid_rule_or_amount_count"], 0)

    def test_trial_balance_run_pins_reconciled_controls_and_source_drillthrough(self):
        run = self.generate_accounting()
        self.assertEqual(run.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertTrue(run.control_gate_required)
        self.assertEqual(run.control_totals["debit"], "1250.00")
        self.assertEqual(run.control_totals["credit"], "1250.00")
        self.assertEqual(run.source_record_count, 1)
        self.assertEqual(len(run.dataset_checksum), 64)
        self.assertEqual(len(run.control_checksum), 64)
        self.assertEqual(len(run.reproduction_key), 64)
        source = run.source_records.get()
        self.assertEqual(source.source_reference, self.entry.reference)
        self.assertEqual(source.snapshot["debit"], "1250.00")
        self.assertEqual(source.snapshot["credit"], "1250.00")
        self.assertEqual(source.source_url, reverse("accounting:entry_detail", args=(self.entry.public_id,)))
        source.amount = Decimal("1.00")
        with self.assertRaisesMessage(ValidationError, "immutable"):
            source.save()

    def test_governed_statement_starters_generate_balanced_explained_statements(self):
        position_mapping = FinanceStatementMapping.objects.get(
            department=self.accounting, statement_type=FinanceStatementMapping.POSITION,
        )
        performance_mapping = FinanceStatementMapping.objects.get(
            department=self.accounting, statement_type=FinanceStatementMapping.PERFORMANCE,
        )
        self.assertEqual(position_mapping.status, FinanceStatementMapping.STARTER)
        self.assertEqual(performance_mapping.status, FinanceStatementMapping.STARTER)
        self.assertEqual(position_mapping.lines.count(), 3)
        self.assertEqual(performance_mapping.lines.count(), 2)

        position_definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_statement_position",
        )
        position = create_manual_run(
            position_definition, position_definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(position.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(position.control_totals["assets"], "1250.00")
        self.assertEqual(position.control_totals["liabilities"], "0.00")
        self.assertEqual(position.control_totals["equity"], "0.00")
        self.assertEqual(position.control_totals["unclosed_operating_result"], "1250.00")
        self.assertEqual(position.control_totals["equation_difference"], "0.00")
        self.assertEqual(position.parameters["_statement_mapping_snapshot"]["version"], 1)
        self.assertEqual(position.parameters["_statement_mapping_checksum"], position_mapping.snapshot_checksum)
        self.assertEqual(position.source_records.get().source_reference, self.entry.reference)

        performance_definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_statement_performance",
        )
        performance = create_manual_run(
            performance_definition, performance_definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(performance.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(performance.control_totals["revenue"], "1250.00")
        self.assertEqual(performance.control_totals["expense"], "0.00")
        self.assertEqual(performance.control_totals["operating_result"], "1250.00")
        self.client.force_login(self.accounting_preparer)
        detail = self.client.get(position_mapping.get_absolute_url())
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "How to maintain a statement mapping")
        self.assertContains(detail, "Coverage passes")
        self.client.force_login(self.budget_preparer)
        self.assertEqual(
            self.client.get(reverse("reporting:statement_mapping_list")).status_code, 403,
        )

    def test_performance_survives_nominal_closing_and_its_reversal(self):
        from accounting.services import create_reversal, submit_entry, post_entry

        self.accounting_preparer.user_permissions.add(Permission.objects.get(
            content_type__app_label="accounting", codename="prepare_journal_entries"))
        self.accounting_reviewer.user_permissions.add(Permission.objects.get(
            content_type__app_label="accounting", codename="post_journal_entries"))
        owner = {"department_id": self.accounting.pk, "department_label": self.accounting.name}
        equity = LedgerAccount.objects.create(
            **owner, code="30101010", title="Accumulated surplus", account_type="equity", normal_balance="credit",
        )
        closing = JournalEntry.objects.create(
            **owner, reference="JEV-CLOSE-001", entry_date=date(2027, 3, 31),
            period=self.period, fund=self.fund, source_type=JournalEntry.CLOSING,
            description="Close nominal revenue into accumulated surplus",
            created_by_id=self.accounting_preparer.pk, created_by_label=self.accounting_preparer.username,
        )
        JournalLine.objects.create(entry=closing, sequence=1, account=self.revenue, debit=Decimal("1250.00"))
        JournalLine.objects.create(entry=closing, sequence=2, account=equity, credit=Decimal("1250.00"))
        submit_entry(closing, self.accounting_preparer)
        post_entry(closing, self.accounting_reviewer)
        definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_statement_performance",
        )
        report = create_manual_run(
            definition, definition.current_template, "xlsx", date(2027, 1, 1), date(2027, 3, 31),
            {}, self.accounting_preparer,
        )
        self.assertEqual(report.control_totals["operating_result"], "1250.00")
        self.assertEqual(report.control_totals["excluded_nominal_closing_entry_count"], 1)
        self.assertEqual(report.source_record_count, 2)
        self.assertTrue(report.source_records.get(source_reference=closing.reference).snapshot["nominal_closing_transfer"])
        from openpyxl import load_workbook
        with report.output_file.open("rb") as stream:
            workbook = load_workbook(stream, data_only=True)
            output_rows = list(workbook.active.values)
        self.assertTrue(any("Surplus / (deficit) for the period" in row and 1250 in row for row in output_rows))
        retained_checksum = report.checksum

        position_definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_statement_position",
        )
        position = create_manual_run(position_definition, position_definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer)
        self.assertEqual(position.control_totals["equity"], "1250.00")
        self.assertEqual(position.control_totals["unclosed_operating_result"], "0.00")
        self.assertEqual(position.control_totals["equation_difference"], "0.00")

        reversal = create_reversal(closing, self.accounting_preparer, reference="JEV-CLOSE-REV",
            entry_date=date(2027, 3, 31), period=self.period, reason="Correct the nominal close")
        submit_entry(reversal, self.accounting_preparer)
        post_entry(reversal, self.accounting_reviewer)
        reopened_report = create_manual_run(definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer)
        self.assertEqual(reopened_report.control_totals["operating_result"], "1250.00")
        self.assertEqual(reopened_report.control_totals["excluded_nominal_closing_entry_count"], 2)
        self.assertEqual(reopened_report.source_record_count, 3)
        report.refresh_from_db()
        self.assertEqual(report.checksum, retained_checksum)

    def test_closing_classification_cannot_hide_a_cash_receipt(self):
        from accounting.services import validate_entry_for_submission
        self.entry.source_type = JournalEntry.CLOSING
        # Classification is checked before submission and again before posting.
        with self.assertRaisesMessage(ValidationError, "cannot contain cash"):
            validate_entry_for_submission(self.entry)

    def _post_statement_adjustment(self, reference, debit_account, credit_account, amount,
                                   *, source_type="adjustment", period=None, entry_date=None, fund=None):
        from accounting.services import submit_entry, post_entry
        self.accounting_preparer.user_permissions.add(Permission.objects.get(
            content_type__app_label="accounting", codename="prepare_journal_entries"))
        self.accounting_reviewer.user_permissions.add(Permission.objects.get(
            content_type__app_label="accounting", codename="post_journal_entries"))
        entry = JournalEntry.objects.create(
            department_id=self.accounting.pk, department_label=self.accounting.name,
            reference=reference, entry_date=entry_date or date(2027, 3, 1), period=period or self.period,
            fund=fund or self.fund, source_type=source_type, description="Synthetic statement adjustment",
            created_by_id=self.accounting_preparer.pk, created_by_label=self.accounting_preparer.username,
        )
        JournalLine.objects.create(entry=entry, sequence=1, account=debit_account, debit=amount)
        JournalLine.objects.create(entry=entry, sequence=2, account=credit_account, credit=amount)
        submit_entry(entry, self.accounting_preparer)
        return post_entry(entry, self.accounting_reviewer)

    def _statement_run(self, dataset_key):
        definition = ReportDefinition.objects.get(department=self.accounting, dataset_key=dataset_key)
        return create_manual_run(definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer)

    def test_contra_assets_reduce_position_instead_of_increasing_it(self):
        owner = {"department_id": self.accounting.pk, "department_label": self.accounting.name}
        allowance = LedgerAccount.objects.create(**owner, code="10301011", title="Allowance for impairment",
            account_type="asset", normal_balance="credit")
        expense = LedgerAccount.objects.create(**owner, code="50501010", title="Impairment loss",
            account_type="expense", normal_balance="debit")
        self._post_statement_adjustment("JEV-ALLOWANCE", expense, allowance, Decimal("100.00"))
        report = self._statement_run("finance_statement_position")
        self.assertEqual(report.control_totals["assets"], "1150.00")
        self.assertEqual(report.control_totals["equation_difference"], "0.00")
        self.assertEqual(report.control_status, ReportRun.CONTROL_RECONCILED)

    def test_contra_revenue_reduces_reported_performance(self):
        refund = LedgerAccount.objects.create(
            department_id=self.accounting.pk, department_label=self.accounting.name,
            code="40101990", title="Revenue refunds", account_type="revenue", normal_balance="debit")
        self._post_statement_adjustment("JEV-REVENUE-REFUND", refund, self.cash, Decimal("50.00"))
        report = self._statement_run("finance_statement_performance")
        self.assertEqual(report.control_totals["revenue"], "1200.00")
        self.assertEqual(report.control_totals["operating_result"], "1200.00")
        self.assertEqual(report.control_status, ReportRun.CONTROL_RECONCILED)

    def _net_assets_opening(self):
        owner = {"department_id": self.accounting.pk, "department_label": self.accounting.name}
        equity = LedgerAccount.objects.create(**owner, code="30101010", title="Accumulated surplus",
            account_type="equity", normal_balance="credit")
        year = FiscalYear.objects.create(**owner, year=2026, label="FY 2026",
            starts_on=date(2026, 1, 1), ends_on=date(2026, 12, 31),
            business_date=date(2026, 3, 31), status=FiscalYear.ACTIVE)
        period = AccountingPeriod.objects.create(**owner, fiscal_year=2026, fiscal_year_record=year,
            period_number=1, label="Prior first quarter", starts_on=date(2026, 1, 1), ends_on=date(2026, 3, 31))
        self._post_statement_adjustment("OPEN-2026", self.cash, equity, Decimal("400.00"),
            source_type="opening", period=period, entry_date=date(2026, 1, 1))
        return equity

    def test_net_assets_export_reconciles_comparison_adjustments_closing_and_reversal(self):
        from accounting.services import create_reversal, submit_entry, post_entry
        from openpyxl import load_workbook
        equity = self._net_assets_opening()
        correction = self._post_statement_adjustment("PRIOR-ERROR", self.cash, equity, Decimal("50.00"),
            source_type=JournalEntry.PRIOR_ERROR)
        self._post_statement_adjustment("DIRECT-REVENUE", self.cash, equity, Decimal("20.00"),
            source_type=JournalEntry.EQUITY_REVENUE)
        self._post_statement_adjustment("CLOSE-CURRENT", self.revenue, equity, Decimal("1250.00"),
            source_type=JournalEntry.CLOSING)
        run = self._statement_run("finance_statement_net_assets")
        self.assertEqual(run.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(run.control_totals["opening"], "400.00")
        self.assertEqual(run.control_totals["closing"], "1720.00")
        self.assertEqual(run.control_totals["operating_result"], "1250.00")
        workbook = load_workbook(run.output_file.path, data_only=True)
        sheet_rows = list(workbook.active.values)
        self.assertTrue(any("prior_error" in row and 50 in row for row in sheet_rows))
        self.assertTrue(any("closing" in row and 1720 in row and 400 in row for row in sheet_rows))
        workbook.close()
        checksum = run.checksum
        reversal = create_reversal(correction, self.accounting_preparer, reference="UNDO-PRIOR",
            entry_date=date(2027, 3, 20), period=self.period, reason="Synthetic corrected prior adjustment")
        submit_entry(reversal, self.accounting_preparer)
        post_entry(reversal, self.accounting_reviewer)
        revised = self._statement_run("finance_statement_net_assets")
        self.assertEqual(revised.control_totals["closing"], "1670.00")
        self.assertEqual(revised.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(revised.source_records.get(source_reference=reversal.reference).snapshot[
            "statement_source_type"], JournalEntry.PRIOR_ERROR)
        run.refresh_from_db()
        self.assertEqual(run.checksum, checksum)

    def test_net_assets_missing_baseline_and_unclassified_equity_block_review(self):
        from accounting.services import create_reversal, submit_entry, post_entry
        missing = self._statement_run("finance_statement_net_assets")
        self.assertEqual(missing.control_status, ReportRun.CONTROL_EXCEPTION)
        equity = self._net_assets_opening()
        entry = self._post_statement_adjustment("UNCLASSIFIED", self.cash, equity, Decimal("10.00"))
        run = self._statement_run("finance_statement_net_assets")
        self.assertEqual(run.control_status, ReportRun.CONTROL_EXCEPTION)
        self.assertEqual(run.control_totals["funds"]["GF"]["current"]["unclassified_entries"], ["UNCLASSIFIED"])
        reversal = create_reversal(entry, self.accounting_preparer, reference="CORRECT-UNCLASSIFIED",
            entry_date=date(2027, 3, 20), period=self.period, reason="Replace with an explicitly classified adjustment")
        submit_entry(reversal, self.accounting_preparer)
        post_entry(reversal, self.accounting_reviewer)
        self._post_statement_adjustment("CLASSIFIED", self.cash, equity, Decimal("10.00"),
            source_type=JournalEntry.PRIOR_ERROR)
        corrected = self._statement_run("finance_statement_net_assets")
        self.assertEqual(corrected.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(corrected.control_totals["closing"], "1660.00")
        self.assertTrue(corrected.source_records.filter(source_reference="UNCLASSIFIED").exists())

    def test_net_assets_definition_cannot_hide_required_rows_or_columns(self):
        definition = ReportDefinition.objects.get(department=self.accounting, dataset_key="finance_statement_net_assets")
        for field, value in (("filters", {"line_code": "closing"}), ("group_by", ["fund_code"]),
                             ("selected_fields", ["amount"])):
            with self.subTest(field=field):
                original = getattr(definition, field)
                setattr(definition, field, value)
                with self.assertRaises(ValueError):
                    build_dataset_with_evidence(definition, date(2027, 1, 1), date(2027, 3, 31), {})
                setattr(definition, field, original)

    def test_direct_equity_classification_cannot_hide_operating_revenue(self):
        from accounting.services import validate_entry_for_submission
        self.entry.source_type = JournalEntry.PRIOR_ERROR
        with self.assertRaisesMessage(ValidationError, "cannot recognize current revenue"):
            validate_entry_for_submission(self.entry)

    def test_net_assets_retains_approved_zero_opening_and_rejects_evidence_drift(self):
        from accounting.models import OpeningBalanceBatch
        from accounting.services import (validate_opening_batch, submit_opening_batch, decide_opening_batch,
            post_opening_batch, reconcile_opening_batch)
        preparer = self.employee(self.accounting, "netassets.opening.preparer", "prepare_opening_balances")
        reviewer = self.employee(self.accounting, "netassets.opening.reviewer",
            "approve_opening_balances", "post_opening_balances")
        preparer.user_permissions.add(Permission.objects.get(content_type__app_label="accounting",
            codename="prepare_opening_balances"))
        reviewer.user_permissions.add(*Permission.objects.filter(content_type__app_label="accounting",
            codename__in=("approve_opening_balances", "post_opening_balances")))
        # Current-only baseline must not invent a prior-year zero comparison.
        batch = OpeningBalanceBatch.objects.create(department_id=self.accounting.pk,
            department_label=self.accounting.name, fiscal_year=self.fiscal_year, period=self.period,
            title="Reviewed zero baseline", source_reference="ZERO-2027", is_zero_balance_declaration=True,
            created_by_id=preparer.pk, created_by_label=preparer.username)
        batch = validate_opening_batch(batch, preparer)
        batch = submit_opening_batch(batch, preparer)
        batch = decide_opening_batch(batch, reviewer, decision=OpeningBalanceBatch.APPROVED,
            evidence_note="Synthetic initial zero balances independently checked")
        batch = post_opening_batch(batch, reviewer)
        batch, _summary = reconcile_opening_batch(batch, reviewer)
        self.assertFalse(batch.postings.exists())
        run = self._statement_run("finance_statement_net_assets")
        current = run.control_totals["funds"]["GF"]["current"]
        self.assertEqual(current["zero_opening_references"], ["ZERO-2027"])
        self.assertEqual(current["movement_difference"], "0.00")
        self.assertEqual(run.control_totals["funds"]["GF"]["comparison"]["opening_baseline_count"], 0)
        self.assertEqual(run.control_status, ReportRun.CONTROL_EXCEPTION)
        self.assertTrue(run.source_records.filter(source_model="OpeningBalanceBatch").exists())
        # Simulate bypassed persistence; no live evidence is changed by this synthetic test.
        OpeningBalanceBatch.objects.filter(pk=batch.pk).update(source_reference="ALTERED-ZERO")
        changed = self._statement_run("finance_statement_net_assets")
        self.assertEqual(changed.control_totals["invalid_zero_openings"], ["ALTERED-ZERO"])

    def test_net_assets_fund_filter_scopes_ledger_controls_and_export_rows_together(self):
        self._net_assets_opening()
        other = Fund.objects.create(department_id=self.accounting.pk, department_label=self.accounting.name,
            code="SEF", name="Special Education Fund")
        self._post_statement_adjustment("SEF-RECEIPT", self.cash, self.revenue, Decimal("90.00"), fund=other)
        definition = ReportDefinition.objects.get(department=self.accounting, dataset_key="finance_statement_net_assets")
        definition.filters = {"fund_code": "GF"}
        _adapter, rows, _totals, evidence = build_dataset_with_evidence(
            definition, date(2027, 1, 1), date(2027, 3, 31), {})
        self.assertEqual({row["fund_code"] for row in rows}, {"GF"})
        self.assertEqual(set(evidence["control_totals"]["funds"]), {"GF"})
        self.assertNotIn("SEF-RECEIPT", [source["source_reference"] for source in evidence["sources"]])
        self.assertEqual(evidence["control_status"], "reconciled")

    def test_net_assets_policy_restatement_and_other_movements_keep_distinct_rows(self):
        equity = self._net_assets_opening()
        for source_type, amount in ((JournalEntry.POLICY_CHANGE, "30.00"),
                                   (JournalEntry.OPENING_RESTATE, "40.00"),
                                   (JournalEntry.EQUITY_OTHER, "60.00")):
            self._post_statement_adjustment(source_type, self.cash, equity, Decimal(amount), source_type=source_type)
        definition = ReportDefinition.objects.get(department=self.accounting, dataset_key="finance_statement_net_assets")
        _adapter, rows, _totals, evidence = build_dataset_with_evidence(
            definition, date(2027, 1, 1), date(2027, 3, 31), {})
        amounts = {row["line_code"]: row["amount"] for row in rows}
        self.assertEqual(amounts["restated_opening"], Decimal("470.00"))
        self.assertEqual(amounts["recognized_result"], Decimal("1250.00"))
        self.assertEqual(amounts["equity_other"], Decimal("60.00"))
        self.assertEqual(amounts["closing"], Decimal("1780.00"))
        self.assertEqual(evidence["control_status"], "reconciled")

    def test_statement_mapping_requires_independent_activation_and_is_immutable(self):
        starter = FinanceStatementMapping.objects.get(
            department=self.accounting, statement_type=FinanceStatementMapping.POSITION,
        )
        successor = FinanceStatementMapping.objects.create(
            department=self.accounting, statement_type=FinanceStatementMapping.POSITION,
            version=2, title="Locally reviewed position mapping", status=FinanceStatementMapping.DRAFT,
            supersedes=starter, authority_reference="Synthetic reviewed COA/GAM statement authority",
            local_acceptance_note="Compared to the signed synthetic local position statement.",
            created_by=self.accounting_preparer,
        )
        FinanceStatementLine.objects.create(
            mapping=successor, position=10, section_code="assets", section_title="Assets",
            line_code="cash", line_title="Cash and cash equivalents",
            selector_type=FinanceStatementLine.ACCOUNT_CODES, account_codes=[self.cash.code],
        )
        submit_statement_mapping(successor, self.accounting_preparer)
        successor.refresh_from_db()
        self.accounting_preparer.user_permissions.add(Permission.objects.get(codename="approve_reports"))
        self.accounting_preparer = get_user_model().objects.get(pk=self.accounting_preparer.pk)
        with self.assertRaisesMessage(ValidationError, "preparer or submitter"):
            review_statement_mapping(successor, self.accounting_preparer, approve=True)
        review_statement_mapping(
            successor, self.accounting_reviewer, approve=True,
            note="Synthetic independent account coverage and signed-reference comparison passed.",
        )
        successor.refresh_from_db()
        starter.refresh_from_db()
        self.assertEqual(successor.status, FinanceStatementMapping.ACTIVE)
        self.assertEqual(starter.status, FinanceStatementMapping.STARTER)
        self.assertEqual(len(successor.snapshot_checksum), 64)
        successor.title = "Silent rewrite"
        with self.assertRaisesMessage(ValidationError, "immutable"):
            successor.save()

        owner = {"department_id": self.accounting.pk, "department_label": self.accounting.name}
        receivable = LedgerAccount.objects.create(
            **owner, code="10301010", title="Receivable", account_type="asset", normal_balance="debit",
        )
        entry = JournalEntry.objects.create(
            **owner, reference="JEV-F9-NEW-ASSET", entry_date=date(2027, 2, 20),
            period=self.period, fund=self.fund, source_type="manual", description="New unmapped asset",
            status=JournalEntry.DRAFT, created_by_id=self.accounting_preparer.pk,
            created_by_label=self.accounting_preparer.username, posted_by_id=self.accounting_reviewer.pk,
            posted_by_label=self.accounting_reviewer.username, posted_at=timezone.now(),
        )
        JournalLine.objects.create(entry=entry, sequence=1, account=receivable, debit=Decimal("25.00"))
        JournalLine.objects.create(entry=entry, sequence=2, account=self.revenue, credit=Decimal("25.00"))
        JournalEntry.objects.filter(pk=entry.pk).update(status=JournalEntry.POSTED)
        definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_statement_position",
        )
        run = create_manual_run(
            definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(run.parameters["_statement_mapping_snapshot"]["version"], 2)
        self.assertEqual(run.control_status, ReportRun.CONTROL_EXCEPTION)
        self.assertEqual(run.control_totals["unmapped_account_codes"], [receivable.code])

    def test_reporting_workspace_explains_latest_statement_measures(self):
        definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_statement_position",
        )
        run = create_manual_run(
            definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.client.force_login(self.accounting_preparer)
        response = self.client.get(reverse("reporting:workspace"))
        self.assertContains(response, "Explained Finance measures")
        self.assertContains(response, "Assets")
        self.assertContains(response, "Source freshness")
        self.assertContains(response, str(run.period_end.year))

    def test_budget_accountability_uses_cumulative_authority_and_posted_movements(self):
        adapter, rows, totals, evidence = build_dataset_with_evidence(
            self.budget_definition, date(2027, 1, 1), date(2027, 3, 31),
            {"_definition_snapshot": {
                "dataset_key": self.budget_definition.dataset_key,
                "selected_fields": self.budget_definition.selected_fields,
                "filters": {}, "group_by": [], "totals": self.budget_definition.totals,
                "sort_by": [],
            }},
        )
        self.assertEqual(adapter.key, "finance_budget_accountability")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["appropriation"], Decimal("100000.00"))
        self.assertEqual(rows[0]["released_allotment"], Decimal("80000.00"))
        self.assertEqual(rows[0]["obligation"], Decimal("50000.00"))
        self.assertEqual(rows[0]["unobligated_allotment"], Decimal("30000.00"))
        self.assertEqual(totals["obligation"], Decimal("50000.00"))
        self.assertEqual(evidence["control_status"], ReportRun.CONTROL_RECONCILED)
        self.assertEqual(len(evidence["sources"]), 3)

    def test_candidate_definition_blocks_official_approval_until_confirmed_successor(self):
        template = self.accounting_definition.current_template
        template.fidelity_status = ReportTemplateVersion.OFFICIAL
        template.fidelity_notes = "Compared with synthetic current form and redacted completed sample."
        template.fidelity_validated_by = self.accounting_reviewer
        template.fidelity_validated_at = timezone.now()
        template.full_clean()
        template.save(update_fields=(
            "fidelity_status", "fidelity_notes", "fidelity_validated_by", "fidelity_validated_at",
        ))
        candidate = self.generate_accounting()
        transition_run(candidate, "review", self.accounting_reviewer, "Controls traced.")
        with self.assertRaisesMessage(ValueError, "local applicability"):
            transition_run(candidate, "approve", self.accounting_reviewer, "Official approval.")
        self.accounting_definition.applicability_status = ReportDefinition.APPLICABILITY_CONFIRMED
        self.accounting_definition.authority_reference = "Synthetic locally reviewed GAM and LGU procedure reference."
        self.accounting_definition.local_acceptance_note = "Confirmed for synthetic UAT by the named Accounting owner."
        self.accounting_definition.full_clean()
        self.accounting_definition.save()
        successor = self.generate_accounting()
        transition_run(successor, "review", self.accounting_reviewer, "Controls traced.")
        transition_run(successor, "approve", self.accounting_reviewer, "Synthetic official acceptance.")
        successor.refresh_from_db()
        self.assertEqual(successor.status, ReportRun.APPROVED)

    def test_control_exception_cannot_enter_review(self):
        broken = JournalEntry.objects.create(
            department_id=self.accounting.pk, department_label=self.accounting.name,
            reference="JEV-F9-BROKEN", entry_date=date(2027, 2, 1), period=self.period,
            fund=self.fund, source_type="manual", description="Synthetic broken posting evidence",
            status=JournalEntry.DRAFT, created_by_id=self.accounting_preparer.pk,
            created_by_label=self.accounting_preparer.username,
            posted_by_id=self.accounting_reviewer.pk, posted_by_label=self.accounting_reviewer.username,
            posted_at=timezone.now(),
        )
        JournalLine.objects.create(entry=broken, sequence=1, account=self.cash, debit=Decimal("10.00"))
        JournalEntry.objects.filter(pk=broken.pk).update(status=JournalEntry.POSTED)
        run = self.generate_accounting()
        self.assertEqual(run.control_status, ReportRun.CONTROL_EXCEPTION)
        with self.assertRaisesMessage(ValueError, "must reconcile"):
            transition_run(run, "review", self.accounting_reviewer)

    def test_control_export_and_reproduction_receipt_are_tracesync_archived(self):
        run = self.generate_accounting()
        self.client.force_login(self.accounting_preparer)
        with tempfile.TemporaryDirectory() as export_root, self.settings(GRAND_EXPORT_ROOT=export_root):
            control = self.client.get(reverse("reporting:run_control_export", args=(run.public_id,)))
            receipt = self.client.get(reverse("reporting:run_reproduction_receipt", args=(run.public_id,)))
            self.assertEqual(control.status_code, 200)
            self.assertEqual(receipt.status_code, 200)
            self.assertEqual(control["X-GRAND-Export-Archived"], "true")
            self.assertEqual(receipt["X-GRAND-Export-Archived"], "true")
            parsed = json.loads(receipt.content)
            self.assertEqual(parsed["checksums"]["dataset_sha256"], run.dataset_checksum)
            self.assertEqual(parsed["checksums"]["reproduction_key"], run.reproduction_key)
            self.assertEqual(parsed["sources"][0]["reference"], self.entry.reference)
            root = Path(export_root)
            self.assertTrue((root / "GRAND_EXPORT_ROOT.json").exists())
            self.assertEqual(len(list(root.rglob("*.manifest.json"))), 2)

    def test_posted_general_ledger_retains_line_level_controls_and_entry_sources(self):
        definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_posted_general_ledger",
        )
        run = create_manual_run(
            definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(run.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(run.row_count, 2)
        self.assertEqual(run.control_totals["debit"], "1250.00")
        self.assertEqual(run.control_totals["credit"], "1250.00")
        self.assertEqual(run.source_record_count, 1)
        source = run.source_records.get()
        self.assertEqual(source.source_reference, self.entry.reference)
        self.assertEqual(len(source.snapshot["lines"]), 2)

    def test_payable_schedule_must_reconcile_to_its_mapped_gl_control(self):
        owner = {"department_id": self.accounting.pk, "department_label": self.accounting.name}
        payable = LedgerAccount.objects.create(
            **owner, code="20101010", title="Accounts payable",
            account_type="liability", normal_balance="credit",
        )
        PostingMapping.objects.create(
            **owner, category=PostingMapping.PAYABLE, source_code="ordinary-supplier",
            label="Ordinary supplier payable", account=payable,
        )
        entry = JournalEntry.objects.create(
            **owner, reference="JEV-F9-AP-0001", entry_date=date(2027, 2, 15),
            period=self.period, fund=self.fund, source_type="voucher",
            source_reference="CASE-F9-AP-1", description="Synthetic payable recognition",
            status=JournalEntry.DRAFT, created_by_id=self.accounting_preparer.pk,
            created_by_label=self.accounting_preparer.username,
            posted_by_id=self.accounting_reviewer.pk,
            posted_by_label=self.accounting_reviewer.username, posted_at=timezone.now(),
        )
        JournalLine.objects.create(
            entry=entry, sequence=1, account=self.cash, debit=Decimal("200.00"),
        )
        payable_line = JournalLine.objects.create(
            entry=entry, sequence=2, account=payable, credit=Decimal("200.00"),
        )
        JournalSubsidiaryLine.objects.create(
            entry=entry, journal_line=payable_line, category=JournalSubsidiaryLine.PAYABLE,
            reference_key="party-f9", reference_label="Synthetic Supplier",
            source_code="ordinary-supplier", source_reference="CASE-F9-AP-1",
            credit=Decimal("200.00"),
        )
        JournalEntry.objects.filter(pk=entry.pk).update(status=JournalEntry.POSTED)
        definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_posted_payable_schedule",
        )
        run = create_manual_run(
            definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(run.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(run.control_totals["subsidiary_balance"], "200.00")
        self.assertEqual(run.control_totals["gl_control_balance"], "200.00")
        self.assertEqual(run.source_records.get().snapshot["reference_key"], "party-f9")

        broken = JournalEntry.objects.create(
            **owner, reference="JEV-F9-AP-BROKEN", entry_date=date(2027, 2, 16),
            period=self.period, fund=self.fund, source_type="manual",
            description="Synthetic control posting without subsidiary detail",
            status=JournalEntry.DRAFT, created_by_id=self.accounting_preparer.pk,
            created_by_label=self.accounting_preparer.username,
            posted_by_id=self.accounting_reviewer.pk,
            posted_by_label=self.accounting_reviewer.username, posted_at=timezone.now(),
        )
        JournalLine.objects.create(entry=broken, sequence=1, account=self.cash, debit=Decimal("50.00"))
        JournalLine.objects.create(entry=broken, sequence=2, account=payable, credit=Decimal("50.00"))
        JournalEntry.objects.filter(pk=broken.pk).update(status=JournalEntry.POSTED)
        exception_run = create_manual_run(
            definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(exception_run.control_status, ReportRun.CONTROL_EXCEPTION)
        self.assertEqual(exception_run.control_totals["difference"], "50.00")
        with self.assertRaisesMessage(ValueError, "must reconcile"):
            transition_run(exception_run, "review", self.accounting_reviewer)

        withholding_definition = ReportDefinition.objects.get(
            department=self.accounting, dataset_key="finance_posted_withholding_schedule",
        )
        withholding_run = create_manual_run(
            withholding_definition, withholding_definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(withholding_run.control_status, ReportRun.CONTROL_EXCEPTION)
        self.assertIn("mapping is not configured", withholding_run.control_message)

    def test_budget_vs_actual_requires_exact_classification_mapping(self):
        owner = {"department_id": self.accounting.pk, "department_label": self.accounting.name}
        center = ResponsibilityCenter.objects.create(
            **owner, code="GSO", name="General Services Office",
            office_id=self.requesting.pk, office_code=self.requesting.slug,
        )
        expense = LedgerAccount.objects.create(
            **owner, code="5-02-03", title="Office operations",
            account_type="expense", normal_balance="debit",
        )
        entry = JournalEntry.objects.create(
            **owner, reference="JEV-F9-ACTUAL-1", entry_date=date(2027, 2, 20),
            period=self.period, fund=self.fund, source_type="voucher",
            source_reference="CASE-F9-ACTUAL-1", description="Synthetic posted actual",
            status=JournalEntry.DRAFT, created_by_id=self.accounting_preparer.pk,
            created_by_label=self.accounting_preparer.username,
            posted_by_id=self.accounting_reviewer.pk,
            posted_by_label=self.accounting_reviewer.username, posted_at=timezone.now(),
        )
        JournalLine.objects.create(
            entry=entry, sequence=1, account=expense, responsibility_center=center,
            debit=Decimal("20000.00"),
        )
        JournalLine.objects.create(entry=entry, sequence=2, account=self.cash, credit=Decimal("20000.00"))
        JournalEntry.objects.filter(pk=entry.pk).update(status=JournalEntry.POSTED)
        definition = ReportDefinition.objects.get(
            department=self.budget, dataset_key="finance_budget_vs_posted_actual",
        )
        adapter, rows, totals, evidence = build_dataset_with_evidence(
            definition, date(2027, 1, 1), date(2027, 3, 31),
            {"_definition_snapshot": {
                "dataset_key": definition.dataset_key,
                "selected_fields": definition.selected_fields, "filters": {}, "group_by": [],
                "totals": definition.totals, "sort_by": [],
            }},
        )
        self.assertEqual(adapter.key, "finance_budget_vs_posted_actual")
        self.assertEqual(rows[0]["posted_actual"], Decimal("20000.00"))
        self.assertEqual(rows[0]["balance_vs_actual"], Decimal("60000.00"))
        self.assertEqual(totals["posted_actual"], Decimal("20000.00"))
        self.assertEqual(evidence["control_status"], ReportRun.CONTROL_RECONCILED)
        self.assertEqual(evidence["control_totals"]["mapping_exception_count"], 0)

        unmatched = LedgerAccount.objects.create(
            **owner, code="5-02-99", title="Unmapped expense",
            account_type="expense", normal_balance="debit",
        )
        broken = JournalEntry.objects.create(
            **owner, reference="JEV-F9-ACTUAL-2", entry_date=date(2027, 2, 21),
            period=self.period, fund=self.fund, source_type="manual",
            description="Synthetic unmatched actual", status=JournalEntry.DRAFT,
            created_by_id=self.accounting_preparer.pk,
            created_by_label=self.accounting_preparer.username,
            posted_by_id=self.accounting_reviewer.pk,
            posted_by_label=self.accounting_reviewer.username, posted_at=timezone.now(),
        )
        JournalLine.objects.create(
            entry=broken, sequence=1, account=unmatched, responsibility_center=center,
            debit=Decimal("100.00"),
        )
        JournalLine.objects.create(entry=broken, sequence=2, account=self.cash, credit=Decimal("100.00"))
        JournalEntry.objects.filter(pk=broken.pk).update(status=JournalEntry.POSTED)
        _adapter, _rows, _totals, exception = build_dataset_with_evidence(
            definition, date(2027, 1, 1), date(2027, 3, 31),
            {"_definition_snapshot": {
                "dataset_key": definition.dataset_key,
                "selected_fields": definition.selected_fields, "filters": {}, "group_by": [],
                "totals": definition.totals, "sort_by": [],
            }},
        )
        self.assertEqual(exception["control_status"], ReportRun.CONTROL_EXCEPTION)
        self.assertEqual(exception["control_totals"]["unmapped_actual"], Decimal("100.00"))

    def test_treasury_register_reconciles_issue_advice_release_and_receipt_evidence(self):
        treasury = Department.objects.create(name="Municipal Treasury Office", slug="f9-treasury")
        treasury_preparer = self.employee(
            treasury, "f9.treasury.preparer",
            "view_reporting_workspace", "generate_reports", "download_reports",
        )
        treasury.deptHead_or_oic = treasury_preparer
        treasury.save(update_fields=("deptHead_or_oic",))
        seed_finance_presets()
        instrument_time = timezone.make_aware(datetime(2027, 3, 5, 10, 30))
        case = VoucherCase.objects.create(
            reference_code="CASE-F9-TRSY-1", transaction_type="ordinary-supplier-claim",
            requesting_department=self.requesting, current_department=treasury,
            payee_name="Synthetic Supplier", particulars="Synthetic released disbursement",
            authoritative_obligation_amount=Decimal("1000.00"),
            current_stage=VoucherCase.COMPLETED, created_by=treasury_preparer,
            completed_at=instrument_time,
        )
        DisbursementVoucher.objects.create(
            case=case, dv_number="DV-F9-TRSY-1", voucher_date=date(2027, 3, 1),
            gross_amount=Decimal("1000.00"), total_deductions=Decimal("100.00"),
            net_amount=Decimal("900.00"), prepared_by=treasury_preparer,
            prepared_at=instrument_time,
        )
        advice = BankAdviceBatch.objects.create(
            advice_number="ADV-F9-1", advice_date=date(2027, 3, 5),
            bank_account_code="GF-CHECKING", status=BankAdviceBatch.ACKNOWLEDGED,
            accounting_department=self.accounting, preparation_note="Synthetic retained advice",
            authority_reference="Synthetic authority", local_applicability_note="Synthetic UAT",
            item_count=1, total_amount=Decimal("900.00"), snapshot_checksum="d" * 64,
            created_by=treasury_preparer, review_submitted_by=treasury_preparer,
            review_submitted_at=instrument_time, approved_by=self.accounting_reviewer,
            approved_at=instrument_time, bank_submitted_by=treasury_preparer,
            bank_submitted_at=instrument_time, submission_reference="SUB-F9-1",
            acknowledged_by=self.accounting_reviewer, acknowledged_at=instrument_time,
            acknowledgement_reference="ACK-F9-1",
        )
        instrument = PaymentInstrument.objects.create(
            case=case, bank_account_code="GF-CHECKING", fund_code="GF",
            check_number="CHK-F9-0001", amount=Decimal("900.00"),
            status=PaymentInstrument.RELEASED, issued_by=treasury_preparer,
            issued_at=instrument_time, released_by=treasury_preparer,
            released_at=instrument_time, released_to="Authorized claimant",
            receipt_reference="RCPT-F9-1", current_advice_batch=advice,
        )
        BankAdviceItem.objects.create(
            batch=advice, instrument=instrument,
            instrument_public_id_snapshot=instrument.public_id,
            check_number_snapshot=instrument.check_number, fund_code_snapshot=instrument.fund_code,
            amount_snapshot=instrument.amount, issued_at_snapshot=instrument.issued_at,
        )
        definition = ReportDefinition.objects.get(
            department=treasury, dataset_key="finance_payment_instrument_register",
        )
        run = create_manual_run(
            definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 12, 31), {}, treasury_preparer,
        )
        self.assertEqual(run.control_status, ReportRun.CONTROL_RECONCILED)
        self.assertEqual(run.control_totals["issued_amount"], "900.00")
        self.assertEqual(run.control_totals["released_amount"], "900.00")
        self.assertEqual(run.source_record_count, 1)
        self.assertEqual(run.source_records.get().source_reference, "CHK-F9-0001")
        self.assertEqual(run.dataset_snapshot["rows"][0]["receipt_reference"], "RCPT-F9-1")

        incomplete_case = VoucherCase.objects.create(
            reference_code="CASE-F9-TRSY-BROKEN", transaction_type="ordinary-supplier-claim",
            requesting_department=self.requesting, current_department=treasury,
            payee_name="Incomplete synthetic payee", particulars="Missing DV source evidence",
            authoritative_obligation_amount=Decimal("10.00"),
            current_stage=VoucherCase.TREASURY_CHECK_PREPARATION, created_by=treasury_preparer,
        )
        PaymentInstrument.objects.create(
            case=incomplete_case, bank_account_code="GF-CHECKING", fund_code="GF",
            check_number="CHK-F9-BROKEN", amount=Decimal("10.00"),
            status=PaymentInstrument.ISSUED, issued_by=treasury_preparer,
            issued_at=instrument_time,
        )
        exception_run = create_manual_run(
            definition, definition.current_template, "xlsx",
            date(2027, 1, 1), date(2027, 12, 31), {}, treasury_preparer,
        )
        self.assertEqual(exception_run.status, ReportRun.GENERATED)
        self.assertEqual(exception_run.control_status, ReportRun.CONTROL_EXCEPTION)
        self.assertEqual(exception_run.control_totals["evidence_exception_count"], 1)

    def test_uat_combined_generator_cannot_create_finance_report(self):
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP

        self.accounting_preparer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        before = ReportRun.objects.count()
        with self.assertRaises(PermissionDenied):
            self.generate_accounting()
        self.assertEqual(ReportRun.objects.count(), before)

    def test_uat_department_head_cannot_review_finance_report(self):
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP

        run = self.generate_accounting()
        before = run.events.count()
        self.accounting_reviewer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        with self.assertRaises(PermissionDenied):
            transition_run(run, "review", self.accounting_reviewer, "Preview must remain read-only")
        run.refresh_from_db()
        self.assertEqual(run.status, ReportRun.GENERATED)
        self.assertEqual(run.events.count(), before)


    def test_finance_manual_generation_rechecks_stored_source_and_actor_authority(self):
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP

        self.accounting_preparer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        self.accounting_definition.dataset_key = "mswd_assistance_volume"
        with self.assertRaises(PermissionDenied):
            self.generate_accounting()
        self.accounting_definition.refresh_from_db()
        with self.assertRaises(PermissionDenied):
            self.generate_accounting(actor=self.budget_preparer)
        self.accounting_reviewer.is_active = False
        with self.assertRaises(PermissionDenied):
            self.generate_accounting(actor=self.accounting_reviewer)
        self.accounting_preparer.groups.clear()
        self.accounting_preparer.user_permissions.clear()
        self.accounting_preparer = get_user_model().objects.get(pk=self.accounting_preparer.pk)
        with self.assertRaises(PermissionDenied):
            self.generate_accounting()
        self.assertFalse(ReportRun.objects.exists())

    def test_finance_uat_run_controls_match_service_and_keep_read_download(self):
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP
        from .run_register_exports import report_action_queryset

        run = self.generate_accounting()
        group = Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0]
        self.accounting_reviewer.groups.add(group)
        self.client.force_login(self.accounting_reviewer)
        detail = self.client.get(reverse("reporting:run_detail", args=(run.public_id,)))
        self.assertEqual(detail.status_code, 200)
        self.assertFalse(detail.context["can_review"])
        self.assertFalse(detail.context["can_approve"])
        self.assertTrue(detail.context["can_download"])
        self.assertEqual(self.client.get(reverse("reporting:run_download", args=(run.public_id,))).status_code, 200)
        self.assertFalse(report_action_queryset(self.accounting_reviewer, "needs_review")[0].exists())
        self.assertEqual(self.client.post(reverse("reporting:run_transition", args=(run.public_id, "review"))).status_code, 403)
        definition = self.client.get(reverse("reporting:definition_detail", args=(run.definition_id,)))
        self.assertFalse(definition.context["can_generate"])
        self.assertEqual(self.client.post(reverse("reporting:definition_detail", args=(run.definition_id,)), {}).status_code, 403)
        self.accounting_reviewer.groups.remove(group)
        transition_run(run, "review", self.accounting_reviewer, "Independent review")
        self.accounting_reviewer.groups.add(group)
        with self.assertRaises(PermissionDenied):
            transition_run(run, "approve", self.accounting_reviewer, "Preview approval")
        run.refresh_from_db()
        self.assertEqual(run.status, ReportRun.REVIEWED)
        self.assertFalse(report_action_queryset(self.accounting_reviewer, "needs_approval")[0].exists())


    def test_statement_mapping_authority_uses_stored_owner_and_denies_uat(self):
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP
        mapping = FinanceStatementMapping.objects.get(
            department=self.accounting, statement_type=FinanceStatementMapping.POSITION,
        )
        mapping.department = self.budget
        with self.assertRaises(PermissionDenied):
            submit_statement_mapping(mapping, self.budget_preparer)
        with self.assertRaises(PermissionDenied):
            review_statement_mapping(mapping, self.budget_preparer, approve=False, note="Foreign return")
        self.accounting_reviewer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        with self.assertRaises(PermissionDenied):
            submit_statement_mapping(mapping, self.accounting_reviewer)
        with self.assertRaises(PermissionDenied):
            review_statement_mapping(mapping, self.accounting_reviewer, approve=False, note="Preview return")
        self.client.force_login(self.accounting_reviewer)
        detail = self.client.get(reverse("reporting:statement_mapping_detail", args=(mapping.public_id,)))
        self.assertEqual(detail.status_code, 200)
        self.assertFalse(detail.context["can_manage"])
        self.assertFalse(detail.context["can_approve"])
        self.assertEqual(self.client.get(reverse("reporting:statement_mapping_create")).status_code, 403)


    def test_uat_cannot_approve_finance_template(self):
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP
        candidate = ReportTemplateVersion.objects.create(
            definition=self.accounting_definition, version=99, title="Preview candidate",
            created_by=self.accounting_preparer, is_active=False,
        )
        self.accounting_reviewer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        self.client.force_login(self.accounting_reviewer)
        response = self.client.post(reverse("reporting:template_approve", args=(candidate.pk,)))
        candidate.refresh_from_db()
        self.assertEqual((response.status_code, candidate.approved_by_id), (403, None))

    def test_uat_cannot_update_finance_definition(self):
        from django.contrib.auth.models import Group
        from django.forms.models import model_to_dict
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP
        from .forms import ReportDefinitionForm
        before = self.accounting_definition.name
        payload = model_to_dict(self.accounting_definition, fields=ReportDefinitionForm.Meta.fields)
        payload["name"] = "Unauthorized preview definition edit"
        self.accounting_preparer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        self.client.force_login(self.accounting_preparer)
        response = self.client.post(reverse("reporting:definition_update", args=(self.accounting_definition.pk,)), payload)
        self.accounting_definition.refresh_from_db()
        self.assertEqual((response.status_code, self.accounting_definition.name), (403, before))

    def test_uat_head_cannot_schedule_finance_output(self):
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP
        from .models import ReportSchedule
        self.accounting_reviewer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        self.client.force_login(self.accounting_reviewer)
        response = self.client.post(reverse("reporting:schedule_create"), {
            "definition": self.accounting_definition.pk,
            "template_version": self.accounting_definition.current_template.pk,
            "name": "Unauthorized Finance schedule", "frequency": ReportSchedule.MONTHLY,
            "output_format": "xlsx", "next_run_at": "2027-04-01T09:00", "is_active": "on",
        })
        self.assertEqual((response.status_code, ReportSchedule.objects.count()), (403, 0))


    def test_uat_finance_template_services_and_controls_recheck_stored_source(self):
        from copy import copy
        from django.contrib.auth.models import Group
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP
        from .models import ReportTemplatePromotion
        from .mappers import preflight_template
        from .template_services import (create_template_promotion, submit_template_promotion,
                                        review_template_promotion, activate_template_promotion, rollback_template_promotion)
        run = self.generate_accounting()
        template = self.accounting_definition.current_template
        promotion = ReportTemplatePromotion.objects.create(
            candidate_template=template, preview_run=run, created_by=self.accounting_preparer,
            change_reason="Synthetic boundary fixture", comparison_note="Synthetic boundary fixture",
            template_checksum="a" * 64,
        )
        self.accounting_reviewer.groups.add(Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0])
        template.definition = copy(template.definition)
        template.definition.dataset_key = "mswd_assistance_volume"
        promotion.candidate_template = template
        calls = (
            lambda: create_template_promotion(template, self.accounting_reviewer, run.period_start, run.period_end,
                                               "xlsx", "Preview change", "Preview comparison"),
            lambda: preflight_template(template, self.accounting_reviewer),
            lambda: submit_template_promotion(promotion, self.accounting_reviewer),
            lambda: review_template_promotion(promotion, self.accounting_reviewer, "return", "Preview return"),
            lambda: activate_template_promotion(promotion, self.accounting_reviewer),
            lambda: rollback_template_promotion(promotion, self.accounting_reviewer, "Preview rollback"),
        )
        for index, call in enumerate(calls):
            with self.subTest(entry_point=index), self.assertRaises(PermissionDenied):
                call()
        promotion.refresh_from_db()
        self.assertEqual(promotion.status, ReportTemplatePromotion.DRAFT)
        self.assertFalse(promotion.events.exists())
        self.client.force_login(self.accounting_reviewer)
        detail = self.client.get(reverse("reporting:template_promotion_detail", args=(promotion.public_id,)))
        self.assertEqual(detail.status_code, 200)
        self.assertFalse(detail.context["can_submit"])
        self.assertFalse(detail.context["can_review"])
        self.assertFalse(detail.context["can_activate"])
        self.assertTrue(detail.context["can_export"])
        definition = self.client.get(reverse("reporting:definition_detail", args=(run.definition_id,)))
        self.assertFalse(definition.context["can_manage_definitions"])
        self.assertFalse(definition.context["can_manage_templates"])
        schedule = self.client.get(reverse("reporting:schedule_create"))
        self.assertFalse(schedule.context["form"].fields["definition"].queryset.filter(pk=run.definition_id).exists())



    def test_manual_generation_pins_stored_template_evidence(self):
        template = self.accounting_definition.current_template
        title = template.title
        template.title = "Unsaved altered title"
        run = create_manual_run(
            self.accounting_definition, template, "xlsx",
            date(2027, 1, 1), date(2027, 3, 31), {}, self.accounting_preparer,
        )
        self.assertEqual(run.parameters["_template_snapshot"]["title"], title)

    def test_generation_ignores_forged_retry_state_on_retained_run(self):
        from unittest.mock import patch
        from .services import generate_report
        run = self.generate_accounting()
        before = (run.status, run.checksum, run.events.count(), run.source_records.count())
        run.status = ReportRun.FAILED
        with patch("reporting.services.build_dataset") as build:
            generate_report(run)
        build.assert_not_called()
        run.refresh_from_db()
        self.assertEqual((run.status, run.checksum, run.events.count(), run.source_records.count()), before)

    def test_schedule_generation_uses_stored_template_and_parameters(self):
        from .models import ReportSchedule
        from .services import execute_schedule
        due = timezone.now()
        template = self.accounting_definition.current_template
        title = template.title
        schedule = ReportSchedule.objects.create(
            definition=self.accounting_definition, template_version=template,
            name="Stored Finance schedule", frequency=ReportSchedule.MONTHLY,
            output_format="xlsx", next_run_at=due, created_by=self.accounting_preparer,
        )
        schedule.template_version.title = "Unsaved scheduled title"
        schedule.parameters = {"unpersisted_parameter": "caller supplied"}
        run, created = execute_schedule(schedule, due)
        self.assertTrue(created)
        self.assertEqual(run.parameters["_template_snapshot"]["title"], title)
        self.assertNotIn("unpersisted_parameter", run.parameters)



    def test_stored_inactive_schedule_cannot_be_reenabled_in_memory(self):
        from .models import ReportSchedule
        from .services import execute_schedule
        schedule = ReportSchedule.objects.create(
            definition=self.accounting_definition,
            template_version=self.accounting_definition.current_template,
            name="Disabled Finance schedule", frequency=ReportSchedule.MONTHLY,
            output_format="xlsx", next_run_at=timezone.now(),
            created_by=self.accounting_preparer, is_active=False,
        )
        schedule.is_active = True
        with self.assertRaisesMessage(ValueError, "Inactive"):
            execute_schedule(schedule)
        self.assertFalse(ReportRun.objects.filter(schedule=schedule).exists())



    def test_personal_report_handoffs_follow_real_review_approval_and_history(self):
        from finance.work_tasks import finance_work_tasks
        template = self.accounting_definition.current_template
        template.fidelity_status = ReportTemplateVersion.OFFICIAL
        template.fidelity_notes = "Synthetic accepted form comparison."
        template.fidelity_validated_by = self.accounting_reviewer
        template.fidelity_validated_at = timezone.now()
        template.save(update_fields=("fidelity_status", "fidelity_notes", "fidelity_validated_by", "fidelity_validated_at"))
        self.accounting_definition.applicability_status = ReportDefinition.APPLICABILITY_CONFIRMED
        self.accounting_definition.authority_reference = "Synthetic LGU authority"
        self.accounting_definition.local_acceptance_note = "Synthetic independent local acceptance."
        self.accounting_definition.save()
        reviewer = self.employee(self.accounting, "report.handoff.reviewer",
                                 "view_reporting_workspace", "review_reports", "view_department_reports")
        run = self.generate_accounting()
        other = self.generate_accounting(self.accounting_reviewer)
        identity = f"report-run:{run.public_id}"
        def rows(actor, view):
            return [row for row in finance_work_tasks(actor, view=view)["tasks"] if row["case_id"] == identity]
        waiting = rows(self.accounting_preparer, "waiting")
        self.assertEqual(len(waiting), 1)
        self.assertIn("Independent report reviewers", waiting[0]["owner_queue"])
        self.assertIsNone(waiting[0]["due_on"])
        self.assertNotIn(f"report-run:{other.public_id}", [row["case_id"] for row in finance_work_tasks(self.accounting_preparer, view="waiting")["tasks"]])
        self.assertEqual(len(rows(self.accounting_preparer, "completed")), 1)
        transition_run(run, "review", reviewer, "Traced exact controls.")
        self.assertIn("Report approvers", rows(self.accounting_preparer, "waiting")[0]["owner_queue"])
        self.assertEqual(len(rows(reviewer, "waiting")), 1)
        self.assertEqual(rows(reviewer, "completed")[0]["subject"], "Reviewed report evidence")
        approval = Permission.objects.get(content_type__app_label="reporting", codename="approve_reports")
        reviewer.user_permissions.add(approval)
        reviewer = get_user_model().objects.get(pk=reviewer.pk)
        self.assertEqual(rows(reviewer, "waiting"), [])
        self.assertEqual(finance_work_tasks(reviewer, view="waiting", display_limit=0)["task_count"], 0)
        reviewer.user_permissions.remove(approval)
        reviewer = get_user_model().objects.get(pk=reviewer.pk)
        self.assertEqual(len(rows(reviewer, "waiting")), 1)
        self.client.force_login(self.accounting_preparer)
        self.assertEqual(self.client.get(waiting[0]["url"]).status_code, 200)
        transition_run(run, "approve", self.accounting_reviewer, "Synthetic acceptance.")
        self.assertEqual(rows(self.accounting_preparer, "waiting"), [])
        self.assertEqual(rows(reviewer, "waiting"), [])
        history = rows(self.accounting_preparer, "completed")
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["source_state"], run.get_status_display())
        self.assertEqual(rows(self.accounting_reviewer, "completed")[0]["subject"], "Approved report output")
        transition_run(run, "supersede", self.accounting_reviewer, "Synthetic successor process.")
        self.assertEqual(len(rows(self.accounting_reviewer, "completed")), 2)
        self.assertEqual(rows(self.accounting_preparer, "completed")[0]["task_id"], history[0]["task_id"])
        self.assertNotEqual(rows(self.accounting_preparer, "completed")[0]["source_version"], history[0]["source_version"])

    def test_report_handoff_views_recheck_current_read_scope_and_preview(self):
        from django.contrib.auth.models import Group
        from finance.work_tasks import finance_work_tasks
        from vouchers.roles import FINANCE_UAT_VIEWER_GROUP
        run = self.generate_accounting()
        identity = f"report-run:{run.public_id}"
        def count(actor, view):
            return sum(row["case_id"] == identity for row in finance_work_tasks(actor, view=view)["tasks"])
        self.assertEqual(count(self.accounting_preparer, "waiting"), 1)
        self.assertEqual(count(self.budget_preparer, "waiting"), 0)
        self.assertEqual(count(self.budget_preparer, "completed"), 0)
        group = Group.objects.get_or_create(name=FINANCE_UAT_VIEWER_GROUP)[0]
        self.accounting_preparer.groups.add(group)
        self.assertEqual(count(self.accounting_preparer, "waiting"), 0)
        self.assertEqual(count(self.accounting_preparer, "completed"), 0)
        self.accounting_preparer.groups.remove(group)
        self.accounting_preparer.user_permissions.remove(Permission.objects.get(
            content_type__app_label="reporting", codename="view_reporting_workspace"))
        actor = get_user_model().objects.get(pk=self.accounting_preparer.pk)
        self.assertEqual(count(actor, "waiting"), 0)
        self.assertEqual(count(actor, "completed"), 0)

    def test_scheduled_generation_is_not_credited_as_manual_completion(self):
        from finance.work_tasks import finance_work_tasks
        from .models import ReportSchedule
        from .services import execute_schedule
        schedule = ReportSchedule.objects.create(
            definition=self.accounting_definition,
            template_version=self.accounting_definition.current_template,
            name="Automatic report", frequency=ReportSchedule.MONTHLY, output_format="xlsx",
            next_run_at=timezone.now(), created_by=self.accounting_preparer,
        )
        run, _created = execute_schedule(schedule)
        self.assertTrue(run.events.filter(action="generated", actor=self.accounting_preparer).exists())
        self.assertNotIn(f"report-run:{run.public_id}", [row["case_id"] for row in finance_work_tasks(self.accounting_preparer, view="completed")["tasks"]])


def tearDownModule():
    shutil.rmtree(FINANCE_REPORT_MEDIA_ROOT, ignore_errors=True)
